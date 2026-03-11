import csv
import io
import json
from datetime import datetime
from pathlib import Path

from docx import Document
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import FileResponse
from openpyxl import Workbook
from pypdf import PdfReader
from reportlab.pdfgen import canvas
from sqlalchemy.orm import Session

from app.ai.ai_business_strategy import generate_business_strategy
from app.ai.ai_chat_assistant import run_assistant
from app.ai.ai_competitor_analysis import generate_competitor_analysis
from app.ai.ai_contract_analyzer import analyze_contract_text
from app.ai.ai_dashboard import generate_dashboard_widgets
from app.ai.ai_market_research import generate_market_research
from app.ai.ai_reports import build_report_content
from app.ai.ai_startup_validator import analyze_startup_idea
from app.core.security import create_access_token, hash_password, verify_password
from app.db import get_db
from app.deps import get_current_user
from app.models import ContractAnalysis, Notification, Report, StartupAnalysis, User
from app.schemas import (
    ChatRequest,
    CompetitorAnalysisRequest,
    LoginRequest,
    MarketResearchRequest,
    SignupRequest,
    StartupIdeaRequest,
    StrategyRequest,
)
from app.services.storage import save_contract_file

router = APIRouter()


def _extract_text_from_file(upload: UploadFile, data: bytes) -> str:
    name = (upload.filename or "").lower()
    if not name:
        raise HTTPException(status_code=400, detail="Uploaded file must have a filename.")
    try:
        if name.endswith(".pdf"):
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        if name.endswith(".docx"):
            doc = Document(io.BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs)
    except Exception:
        raise HTTPException(status_code=400, detail="Unable to read file. Upload a valid PDF or DOCX.")
    raise HTTPException(status_code=400, detail="Unsupported file type. Use PDF or DOCX.")


def _wrap_text(text: str, max_chars: int = 95) -> list[str]:
    words = (text or "").split()
    if not words:
        return [""]
    lines: list[str] = []
    current = words[0]
    for word in words[1:]:
        candidate = f"{current} {word}"
        if len(candidate) <= max_chars:
            current = candidate
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


@router.post("/auth/signup")
def signup(payload: SignupRequest, db: Session = Depends(get_db)):
    exists = db.query(User).filter(User.email == payload.email).first()
    if exists:
        raise HTTPException(status_code=400, detail="Email already exists")

    user = User(name=payload.name, email=payload.email, password_hash=hash_password(payload.password))
    db.add(user)
    db.commit()
    db.refresh(user)
    token = create_access_token({"sub": str(user.id)})
    return {"access_token": token, "user": {"name": user.name, "email": user.email}}


@router.post("/auth/login")
def login(payload: LoginRequest, db: Session = Depends(get_db)):
    user = db.query(User).filter(User.email == payload.email).first()
    if not user or not verify_password(payload.password, user.password_hash):
        raise HTTPException(status_code=401, detail="Invalid credentials")

    token = create_access_token({"sub": str(user.id)})
    return {"access_token": token, "user": {"name": user.name, "email": user.email}}


@router.get("/dashboard/summary")
def dashboard_summary(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    startup = db.query(StartupAnalysis).filter(StartupAnalysis.user_id == current_user.id).order_by(StartupAnalysis.id.desc()).first()
    contract = db.query(ContractAnalysis).filter(ContractAnalysis.user_id == current_user.id).order_by(ContractAnalysis.id.desc()).first()

    success_score = startup.success_probability if startup else 74
    contract_risk = contract.risk_score if contract else 42
    market_demand = startup.market_demand if startup else 81
    competitor_count = int((startup.competition_level if startup else 58) // 8)

    trend = [{"month": m, "success": v} for m, v in zip(["Jan", "Feb", "Mar", "Apr", "May", "Jun"], [62, 67, 71, 73, 76, success_score])]
    demand = [{"month": m, "demand": v} for m, v in zip(["Jan", "Feb", "Mar", "Apr", "May", "Jun"], [58, 61, 69, 74, 79, market_demand])]
    risk_dist = [
        {"name": "Low", "value": max(10, 100 - contract_risk - 20)},
        {"name": "Medium", "value": 20},
        {"name": "High", "value": contract_risk},
    ]
    competitor_growth = [{"q": q, "count": c} for q, c in zip(["Q1", "Q2", "Q3", "Q4"], [4, 6, 7, competitor_count])]

    notifications = db.query(Notification).order_by(Notification.id.desc()).limit(5).all()
    widgets = generate_dashboard_widgets(startup, contract, notifications)

    return {
        "cards": {
            "startup_success_score": round(success_score, 1),
            "contract_risk_score": round(contract_risk, 1),
            "market_demand_score": round(market_demand, 1),
            "competitor_count": competitor_count,
        },
        "charts": {
            "startup_trend": trend,
            "market_demand": demand,
            "contract_risk_distribution": risk_dist,
            "competitor_growth": competitor_growth,
        },
        "widgets": {
            "weekly_insights": widgets["weekly_insights"],
            "recent_analyses": widgets["recent_analyses"],
            "recommendations": widgets["recommendations"],
            "notifications": [
                {"title": n.title, "message": n.message, "category": n.category} for n in notifications
            ],
        },
    }


@router.post("/startup-idea-validator/analyze")
def startup_idea_validator(
    payload: StartupIdeaRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    result = analyze_startup_idea(payload.idea, payload.target_audience, payload.industry, payload.revenue_model)

    row = StartupAnalysis(
        user_id=current_user.id,
        idea=payload.idea,
        audience=payload.target_audience,
        industry=payload.industry,
        revenue_model=payload.revenue_model,
        success_probability=result["success_probability"],
        market_demand=result["market_demand"],
        competition_level=result["competition_level"],
        profit_potential=result["profit_potential"],
    )
    db.add(row)
    db.commit()
    return result


@router.post("/contract-risk-analyzer/analyze")
async def contract_risk_analyzer(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    content = await file.read()
    path_or_url = save_contract_file(file.filename, content)
    text = _extract_text_from_file(file, content)
    result = analyze_contract_text(text)

    row = ContractAnalysis(
        user_id=current_user.id,
        file_name=file.filename,
        s3_url=path_or_url,
        summary=str(result["summary"]),
        risk_score=float(result["risk_score"]),
        risky_clauses=json.dumps(result["risky_clauses"]),
    )
    db.add(row)
    db.commit()

    return {"file_url": path_or_url, **result}


@router.post("/market-research/generate")
def market_research(payload: MarketResearchRequest, current_user: User = Depends(get_current_user)):
    return generate_market_research(payload.industry)


@router.post("/competitor-analysis/generate")
def competitor_analysis(payload: CompetitorAnalysisRequest, current_user: User = Depends(get_current_user)):
    return generate_competitor_analysis(payload.idea)


@router.post("/business-strategy/generate")
def business_strategy(payload: StrategyRequest, current_user: User = Depends(get_current_user)):
    return generate_business_strategy(payload.idea)


@router.post("/assistant/chat")
def assistant_chat(payload: ChatRequest, current_user: User = Depends(get_current_user)):
    return run_assistant(payload.question, payload.history)


@router.get("/notifications")
def list_notifications(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    rows = db.query(Notification).order_by(Notification.id.desc()).limit(20).all()
    return [{"title": r.title, "message": r.message, "category": r.category, "created_at": str(r.created_at)} for r in rows]


@router.get("/reports/download/{report_type}/{file_format}")
def download_report(
    report_type: str,
    file_format: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    reports_dir = Path("reports")
    reports_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")

    startup = db.query(StartupAnalysis).filter(StartupAnalysis.user_id == current_user.id).order_by(StartupAnalysis.id.desc()).first()
    contract = db.query(ContractAnalysis).filter(ContractAnalysis.user_id == current_user.id).order_by(ContractAnalysis.id.desc()).first()
    report_content = build_report_content(report_type, current_user.email, startup, contract)

    content_rows = [
        ["metric", "value"],
        ["report_type", report_type],
        ["generated_for", current_user.email],
        ["generated_at", timestamp],
        ["title", report_content["title"]],
        ["summary", report_content["summary"]],
    ]
    for index, item in enumerate(report_content.get("highlights", []), start=1):
        content_rows.append([f"highlight_{index}", item])
    for item in report_content.get("metrics", []):
        content_rows.append([str(item.get("metric") or "metric"), str(item.get("value") or "")])

    if file_format == "csv":
        path = reports_dir / f"{report_type}_{timestamp}.csv"
        with path.open("w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerows(content_rows)
        file_path = str(path)

    elif file_format == "excel":
        path = reports_dir / f"{report_type}_{timestamp}.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Report"
        for row in content_rows:
            ws.append(row)
        wb.save(path)
        file_path = str(path)

    elif file_format == "pdf":
        path = reports_dir / f"{report_type}_{timestamp}.pdf"
        c = canvas.Canvas(str(path))
        text = c.beginText(72, 800)
        text.setFont("Helvetica-Bold", 16)
        text.textLine(report_content["title"])
        text.moveCursor(0, 12)
        text.setFont("Helvetica", 11)
        for line in _wrap_text(report_content["summary"]):
            text.textLine(line)
        text.moveCursor(0, 12)
        text.setFont("Helvetica-Bold", 12)
        text.textLine("Highlights")
        text.setFont("Helvetica", 11)
        for item in report_content.get("highlights", []):
            for line in _wrap_text(f"- {item}"):
                text.textLine(line)
        text.moveCursor(0, 12)
        text.setFont("Helvetica-Bold", 12)
        text.textLine("Metrics")
        text.setFont("Helvetica", 11)
        for item in report_content.get("metrics", []):
            metric = str(item.get("metric") or "metric").replace("_", " ").title()
            value = str(item.get("value") or "")
            for line in _wrap_text(f"- {metric}: {value}"):
                text.textLine(line)
        c.drawText(text)
        c.save()
        file_path = str(path)

    else:
        raise HTTPException(status_code=400, detail="Unsupported format")

    db.add(Report(user_id=current_user.id, report_type=report_type, file_format=file_format, file_path=file_path))
    db.commit()
    return FileResponse(file_path, filename=Path(file_path).name)


@router.get("/profile")
def profile(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    startups = db.query(StartupAnalysis).filter(StartupAnalysis.user_id == current_user.id).order_by(StartupAnalysis.id.desc()).limit(5).all()
    contracts = db.query(ContractAnalysis).filter(ContractAnalysis.user_id == current_user.id).order_by(ContractAnalysis.id.desc()).limit(5).all()
    reports = db.query(Report).filter(Report.user_id == current_user.id).order_by(Report.id.desc()).limit(5).all()

    return {
        "user": {"name": current_user.name, "email": current_user.email, "created_at": str(current_user.created_at)},
        "saved_startup_ideas": [{"idea": s.idea, "score": s.success_probability, "created_at": str(s.created_at)} for s in startups],
        "uploaded_contracts": [{"file": c.file_name, "risk_score": c.risk_score, "created_at": str(c.created_at)} for c in contracts],
        "generated_reports": [{"type": r.report_type, "format": r.file_format, "path": r.file_path, "created_at": str(r.created_at)} for r in reports],
    }